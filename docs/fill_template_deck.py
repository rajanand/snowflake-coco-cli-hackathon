"""Fills slides 3, 4, 5 of the official hackathon submission template with
content sourced from this repo, per the guidelines listed on slide 2
(Problem Brief / Architecture Diagram / Impact Statement).

Edits the template in place:
  "docs/Prototype Submission Template _ CoCo CLI Hackathon GCC Edition.pptx"

Every shape added lives in the transparent white canvas between the black
header bar (ends ~0.53in) and the blue gradient footer bar (starts ~5.53in)
on the template's 10in x 5.625in slides -- measured by scanning the alpha
channel of each slide's background picture.

Run: python docs/fill_template_deck.py
"""

from __future__ import annotations

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

TEMPLATE_PATH = (
    "docs/Prototype Submission Template _ CoCo CLI Hackathon GCC Edition.pptx"
)

DARK = RGBColor(0x20, 0x27, 0x29)      # matches slide 1's label color
MUTED = RGBColor(0x5B, 0x64, 0x66)
ACCENT = RGBColor(0x00, 0x74, 0xD6)    # sampled from the template's footer gradient
LINE = RGBColor(0xDD, 0xE3, 0xE5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Manrope"                        # the template's own font; degrades gracefully

prs = Presentation(TEMPLATE_PATH)


def textbox(slide, left, top, width, height):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    return tb, tf


def set_run(run, text, size, color, bold=False, italic=False):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = FONT


def heading(slide, text):
    _, tf = textbox(slide, Inches(0.45), Inches(0.6), Inches(9.1), Inches(0.5))
    p = tf.paragraphs[0]
    r = p.add_run()
    set_run(r, text, 24, ACCENT, bold=True)
    rule = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(0.47), Inches(1.14), Inches(2.0), Inches(1.14)
    )
    rule.line.color.rgb = ACCENT
    rule.line.width = Pt(2.25)


def label_answer_block(slide, left, top, width, height, label, answer, label_size=12.5, ans_size=10.8):
    _, tf = textbox(slide, left, top, width, Inches(0.3))
    p = tf.paragraphs[0]
    r = p.add_run()
    set_run(r, label, label_size, ACCENT, bold=True)
    _, tf2 = textbox(slide, left, top + Inches(0.32), width, height - Inches(0.32))
    p2 = tf2.paragraphs[0]
    p2.line_spacing = 1.12
    r2 = p2.add_run()
    set_run(r2, answer, ans_size, DARK)


def column_block(slide, left, top, width, title, bullets, title_size=13.5, bullet_size=10.6):
    _, tf = textbox(slide, left, top, width, Inches(0.35))
    p = tf.paragraphs[0]
    r = p.add_run()
    set_run(r, title, title_size, ACCENT, bold=True)
    y = top + Inches(0.4)
    for b in bullets:
        _, tfb = textbox(slide, left, y, width, Inches(0.9))
        pb = tfb.paragraphs[0]
        pb.line_spacing = 1.1
        rb = pb.add_run()
        set_run(rb, "\u2014  " + b, bullet_size, DARK)
        # rough height estimate based on text length vs column width
        chars_per_line = max(1, int(width / Inches(0.082)))
        n_lines = max(1, -(-len(b) // chars_per_line))
        y += Inches(0.02) + Inches(0.21) * n_lines


# ---------------------------------------------------------------------------
# Slide 3 (index 2) -- Problem Brief
# source: README.md problem table + sql/00_source_schemas/README.md +
#         sql/02_silver/README.md + sql/04_semantic_view/README.md
# ---------------------------------------------------------------------------

def fill_slide_problem_brief():
    s = prs.slides[2]
    heading(s, "Problem Brief")

    qa = [
        ("What real business problem does this solve?",
         "Supply chain data is scattered across ERP, TMS, Supplier Portal, IoT sensor, and "
         "Finance systems, each computing the same KPI differently. Ask \u201cwhat's our "
         "On-Time Delivery rate?\u201d and four systems give four different, individually "
         "defensible answers (89.4% / 91.3% / 78.6% / a biased IoT sample) \u2014 there is no "
         "shared source of truth."),
        ("Who is the target user / persona?",
         "Planning, Procurement, and Logistics analysts inside a Global Capability Center who "
         "all report the same supply-chain KPIs upward, but pull them from different source "
         "systems using different, undocumented definitions."),
        ("Current pain point \u2192 how this improves it",
         "Today each team's number is \u201ccorrect\u201d by its own definition, so leadership can't "
         "trust cross-team reporting. This project encodes ONE governed definition per metric "
         "as a native Snowflake Semantic View \u2014 every persona now gets the identical number "
         "(64.875% OTD), verified regardless of how the question is worded."),
        ("Industry / domain context",
         "Manufacturing / industrial supply chain \u2014 Supplier \u2192 Part \u2192 Plant \u2192 Shipment \u2192 "
         "Order \u2192 Customer, the exact entity chain called out in the challenge brief. 100% "
         "synthetic data; no proprietary or real GCC data used anywhere in the project."),
    ]
    col_w = Inches(4.35)
    row_h = Inches(1.72)
    left0 = Inches(0.45)
    top0 = Inches(1.55)
    gap_x = Inches(0.3)
    for i, (label, answer) in enumerate(qa):
        r, c = divmod(i, 2)
        left = left0 + c * (col_w + gap_x)
        top = top0 + r * row_h
        label_answer_block(s, left, top, col_w, row_h, label, answer)


# ---------------------------------------------------------------------------
# Slide 4 (index 3) -- Architecture Diagram
# source: README.md Architecture section, docs/architecture-and-er-diagram.html,
#         agents/agent_router.py, .snowflake/cortex/plans + .cortex/plans file list
# ---------------------------------------------------------------------------

def fill_slide_architecture():
    s = prs.slides[3]
    heading(s, "Architecture Diagram")

    stages = [
        ("Source\nSystems", "5 fragmented\nSOURCE_* schemas"),
        ("Bronze", "Unified raw\nlanding zone"),
        ("Silver", "Entity resolution\n(shipment_crosswalk)"),
        ("Gold", "9-entity model,\n28 Dynamic Tables"),
        ("Semantic\nView", "Governed metrics\n+ synonyms"),
        ("Agents\n+ App", "Multi-agent router,\nStreamlit app"),
    ]
    n = len(stages)
    top = Inches(1.5)
    box_w = Inches(1.32)
    box_h = Inches(1.05)
    gap = Inches(0.19)
    x = Inches(0.45)
    centers = []
    for title, sub in stages:
        box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, top, box_w, box_h)
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(0xF2, 0xF7, 0xFC)
        box.line.color.rgb = ACCENT
        box.line.width = Pt(1.1)
        box.shadow.inherit = False
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.05)
        tf.margin_right = Inches(0.05)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        first = True
        for line in title.split("\n"):
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            set_run(r, line, 10.5, DARK, bold=True)
        for line in sub.split("\n"):
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.CENTER
            r2 = p2.add_run()
            set_run(r2, line, 8, MUTED)
        centers.append((x, x + box_w))
        x += box_w + gap

    mid_y = top + box_h // 2
    for i in range(n - 1):
        arrow = s.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, centers[i][1], mid_y, centers[i + 1][0], mid_y
        )
        arrow.line.color.rgb = ACCENT
        arrow.line.width = Pt(1.25)
        line_elem = arrow.line._get_or_add_ln()
        tail = line_elem.makeelement(
            qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}
        )
        line_elem.append(tail)

    cols = [
        ("CoCo CLI skills used",
         "Built end-to-end with CoCo CLI across 8 dated planning sessions "
         "(2026-09-12 \u2192 09-23, committed under .snowflake/cortex/plans/ and .cortex/plans/): "
         "sql-author for the medallion SQL, agent-studio for the Semantic View, "
         "developing-with-streamlit-in-snowflake for the app, and verification skills for the "
         "golden-question test suite."),
        ("Data sources",
         "Structured: ERP orders, TMS deliveries, Supplier Portal shipments/POs, IoT sensor "
         "tracking events, Finance COGS/overhead. Unstructured: 11 synthetic supplier contract "
         "PDFs, parsed via AI_PARSE_DOCUMENT and indexed in Cortex Search."),
        ("How components plug together",
         "Each medallion layer is its own set of Dynamic Tables (TARGET_LAG-scheduled). The "
         "Semantic View is the single governed interface on top; the agent router (4 sub-agents "
         "behind a 5-intent classifier) and the Streamlit app both consume only that interface."),
    ]
    col_w = Inches(2.93)
    left0 = Inches(0.45)
    top0 = Inches(2.85)
    gap_x = Inches(0.14)
    for i, (title, body) in enumerate(cols):
        left = left0 + i * (col_w + gap_x)
        label_answer_block(s, left, top0, col_w, Inches(2.35), title, body,
                            label_size=11.5, ans_size=9.6)


# ---------------------------------------------------------------------------
# Slide 5 (index 4) -- Impact Statement
# source: sql/02_silver/README.md, sql/04_semantic_view/README.md,
#         sql/07_rbac/README.md, sql/08_marketplace_enrichment/README.md,
#         docs/pending-gaps.md
# ---------------------------------------------------------------------------

def fill_slide_impact():
    s = prs.slides[4]
    heading(s, "Impact Statement")

    # retitle the existing bottom tag from "Additional Slide" to match this section
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip() == "Additional Slide":
            sh.text_frame.paragraphs[0].runs[0].text = "Impact Statement"

    cols = [
        ("Measurable outcomes", [
            "4 conflicting legacy answers (89.4% / 91.3% / 78.6% / biased) collapse into "
            "1 governed number: 64.875%.",
            "100% entity resolution across 4 independent systems (800/800 shipments matched).",
            "Cross-persona proof: 3 differently-worded questions \u2192 identical value + identical "
            "base_table/measure_name.",
        ]),
        ("Scalability potential", [
            "28 Dynamic Tables, TARGET_LAG-scheduled \u2014 the pipeline scales with data volume, "
            "no manual reconciliation.",
            "Least-privilege RBAC split (SUPPLY_CHAIN_ANALYST_RO) is negative-tested and ready "
            "for enterprise rollout.",
            "The Semantic View pattern extends to new source systems or metrics without "
            "touching the agents or the app.",
        ]),
        ("Beyond the demo", [
            "The multi-agent pattern (what-if simulation, root-cause trace, document Q&A) is "
            "domain-agnostic \u2014 reusable for any governed-analytics use case.",
            "Marketplace enrichment (free Pelmorex weather data) proves the pattern for pulling "
            "in third-party signals.",
            "Disclosed roadmap (MCP connectors, Task automation, a packaged CoCo skill) is the "
            "explicit productionization path, not a hidden gap.",
        ]),
    ]
    col_w = Inches(2.93)
    left0 = Inches(0.45)
    top0 = Inches(1.55)
    gap_x = Inches(0.14)
    for i, (title, bullets) in enumerate(cols):
        left = left0 + i * (col_w + gap_x)
        column_block(s, left, top0, col_w, title, bullets)


def build():
    fill_slide_problem_brief()
    fill_slide_architecture()
    fill_slide_impact()
    prs.save(TEMPLATE_PATH)
    print(f"Updated slides 3-5 in: {TEMPLATE_PATH}")


if __name__ == "__main__":
    build()
